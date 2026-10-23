"""Число строк учеников в Excel-отчёте следует за списком, блок итогов остаётся под таблицей."""
import json
from pathlib import Path

from openpyxl import load_workbook

from mektep_core.build_report import build_report, summary_rows
from mektep_core.build_word_report import (
    _extract_names_from_excel_column,
    _extract_sheet_block,
)

REPO = Path(__file__).resolve().parent.parent
TEMPLATE = REPO / "Шаблон.xlsx"


def _write_inputs(tmp_path: Path, count: int) -> tuple[Path, Path]:
    # 10 → «5», 9 → «5», 7 → «4», 5 → «3», 2 → «2» при максимуме 10.
    scores = [10, 9, 7, 5, 2]
    students = [
        {
            "num": i + 1,
            "fio": f"Ученик {i + 1}",
            "quarter_num": 2,
            "average": scores[i % len(scores)],
            "grade": "",
            "points": {},
        }
        for i in range(count)
    ]
    students_path = tmp_path / "students.json"
    context_path = tmp_path / "context.json"
    students_path.write_text(json.dumps(students, ensure_ascii=False), encoding="utf-8")
    context_path.write_text(
        json.dumps(
            {
                "class": "5 Е",
                "subject": "Математика",
                "org_name": "Лицей",
                "profile_name": "Учитель",
                "has_quarter_grade_header": False,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return students_path, context_path


def _formative_sheet(tmp_path: Path, count: int):
    students_path, context_path = _write_inputs(tmp_path, count)
    out = build_report(TEMPLATE, students_path, context_path, tmp_path / "out")
    wb = load_workbook(out)
    return wb["Формативное оценивание"]


def test_fewer_students_shrink_the_table(tmp_path):
    ws = _formative_sheet(tmp_path, 5)
    assert ws["B8"].value == "Ученик 1"
    assert ws["B12"].value == "Ученик 5"
    assert ws["B13"].value in (None, "")
    assert ws["A12"].value == 5
    totals = summary_rows(ws)
    assert totals["5"] == 14
    assert str(ws[f"D{totals['5']}"].value).startswith("Всего")
    assert ws[f"F{totals['5']}"].value == 2
    assert ws["J8"].value == 60
    assert ws["B7"].value == "ФИО"
    assert ws["B1"].value == "Лицей"
    assert any(
        merged.min_row == totals["5"] and merged.min_col == 4
        for merged in ws.merged_cells.ranges
    )
    assert ws["B12"].border.left.style == "thin"


def test_more_than_template_capacity_keeps_every_student(tmp_path):
    ws = _formative_sheet(tmp_path, 40)
    assert ws["B47"].value == "Ученик 40"
    assert ws["A47"].value == 40
    assert ws["B48"].value in (None, "")
    totals = summary_rows(ws)
    assert totals["5"] == 49
    assert ws[f"F{totals['5']}"].value == 16
    assert ws[f"G{totals['4']}"].value == 8
    assert ws[f"H{totals['3']}"].value == 8
    assert ws[f"I{totals['2']}"].value == 8
    assert ws["C4"].value == 40
    assert ws["J8"].value == 60
    assert ws["B47"].border.left.style == "thin"
    assert "B1:N1" in {str(merged) for merged in ws.merged_cells.ranges}
    assert any(
        merged.min_row == totals["5"] and merged.min_col == 4
        for merged in ws.merged_cells.ranges
    )

    block = _extract_sheet_block(ws)
    assert block["count_5"] == 16
    assert block["students"] == 40
    high = _extract_names_from_excel_column(ws, "L")
    low = _extract_names_from_excel_column(ws, "N")
    assert "Ученик 37" in high
    assert "Ученик 40" in low
    assert len(high) == 16


def test_exact_template_capacity_keeps_summary_in_place(tmp_path):
    ws = _formative_sheet(tmp_path, 32)
    assert ws["B39"].value == "Ученик 32"
    assert summary_rows(ws)["5"] == 41
    assert ws["F41"].value == 14
