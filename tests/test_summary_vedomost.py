"""Сводная ведомость: колонки бланка 5д и скачивание из кабинета."""

from __future__ import annotations

import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from webapp import create_app
from webapp.config import TestingConfig
from webapp.extensions import db
from webapp.models import Class, GradeReport, Role, School, User
from webapp.services.academic_year import current_academic_year
from webapp.services.grade_reports.excel.summary_vedomost import (
    TEMPLATE_PATH,
    build_summary_vedomost_workbook,
)


class _Report:
    def __init__(self, subject, period_type, period_number, students):
        self.subject_name = subject
        self.period_type = period_type
        self.period_number = period_number
        self.grades_json = json.dumps({"students": students}, ensure_ascii=False)


def _grade(name, grade):
    return {"name": name, "grade": grade}


def _workbook(reports):
    output, filename = build_summary_vedomost_workbook(
        class_name="5Д",
        academic_year=2025,
        school_name="Специализированный IT лицей",
        class_teacher="Баер Эдуард Викторович",
        reports=reports,
        director="",
    )
    assert filename == "Сводная_ведомость_5Д_2025.xlsx"
    return load_workbook(output).active


def test_template_has_no_student_names():
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    assert ws.max_row == 10
    blob = " ".join(
        str(cell.value) for row in ws.iter_rows() for cell in row if cell.value
    )
    assert "Амангельдиева" not in blob
    assert "Шетел тілі" in blob


def test_tracks_duplicates_year_and_elective():
    reports = [
        _Report(
            "Информатика Информатика (2)",
            "quarter",
            2,
            [_grade("Амангельдиева Гаухар", 5)],
        ),
        _Report(
            "Иностранный язык Иностранный язык (1)",
            "quarter",
            2,
            [_grade("Ли Арина", 4)],
        ),
        _Report(
            "Иностранный язык Иностранный язык (2)",
            "quarter",
            1,
            [_grade("Амангельдиева Гаухар", 5)],
        ),
        _Report("Технология (1)", "quarter", 1, [_grade("Ли Арина", 4)]),
        _Report("Дене шынықтыру", "quarter", 1, [_grade("Ли Арина", "зачтено")]),
        _Report("Математика", "quarter", 1, [_grade("Ли Арина", 5)]),
        _Report("Математика", "quarter", 2, [_grade("Ли Арина", 5)]),
        _Report("Математика", "quarter", 3, [_grade("Ли Арина", 4)]),
        _Report("Математика", "quarter", 4, [_grade("Ли Арина", 4)]),
        _Report("Математика", "final", 1, [_grade("Ли Арина", 3)]),
        _Report("Химия", "quarter", 1, [_grade("Ли Арина", 5)]),
    ]
    ws = _workbook(reports)

    assert ws["A5"].value == "Специализированный IT лицей"
    assert ws["A8"].value == "5 сынып / класс  2025  оқу жылы / учебный год"
    assert ws["A11"].value == 1
    assert ws["B11"].value == "Амангельдиева Гаухар"
    assert ws["D11"].value == 5  # шетел тілі, казахское отделение, 1 четверть
    assert ws["F11"].value == " "
    assert ws["I12"].value == 5  # информатика, казахское отделение
    assert ws["K12"].value == " "

    assert ws["B18"].value == "Ли Арина"
    assert ws["F19"].value == 4  # шетел тілі, русское отделение
    assert ws["D19"].value == " "
    assert ws["G18"].value == 4  # технология в обе колонки
    assert ws["V18"].value == 4
    assert ws["T18"].value == "зч"
    assert ws["U18"].value == "зч"
    assert ws["L22"].value == 5  # годовая: (5+5+4+4)/4
    assert ws["L23"].value == " "  # экзамен пустой
    assert ws["L24"].value == 3  # итог
    assert ws["X10"].value == "Химия"
    assert ws["X18"].value == 5
    assert ws["B27"].value == "Орта білім беру ұйымының директоры"
    assert ws["V31"].value == "Баер Эдуард Викторович"


def test_semester_replaces_odd_quarters():
    reports = [
        _Report("Всемирная история", "quarter", 1, [_grade("Ли Арина", 2)]),
        _Report("Всемирная история", "semester", 1, [_grade("Ли Арина", 5)]),
    ]
    ws = _workbook(reports)
    assert ws["S11"].value == " "  # 1 четверть снята: предмет полугодовой
    assert ws["S12"].value == 5


@pytest.fixture
def app():
    application = create_app(TestingConfig)
    with application.app_context():
        db.create_all()
        yield application
        db.session.remove()
        db.drop_all()


@pytest.fixture
def client(app):
    return app.test_client()


def test_admin_page_downloads_vedomost(app, client):
    with app.app_context():
        year = current_academic_year()
        school = School(name="Лицей ведомости", is_active=True)
        db.session.add(school)
        db.session.flush()
        admin = User(
            username="vedomost-admin",
            full_name="Админ А.",
            role=Role.SCHOOL_ADMIN.value,
            school_id=school.id,
            is_active=True,
        )
        admin.set_password("secret123")
        teacher = User(
            username="vedomost-teacher",
            full_name="Классный Руководитель",
            role=Role.TEACHER.value,
            school_id=school.id,
            is_active=True,
        )
        teacher.set_password("secret123")
        db.session.add_all([admin, teacher])
        db.session.flush()
        db.session.add(Class(school_id=school.id, name="5Д", class_teacher_id=teacher.id))
        db.session.add(
            GradeReport(
                teacher_id=teacher.id,
                school_id=school.id,
                class_name="5Д",
                subject_name="Математика",
                period_type="quarter",
                period_number=2,
                academic_year=year,
                grades_json=json.dumps(
                    {"students": [_grade("Ли Арина", 4)]},
                    ensure_ascii=False,
                ),
            )
        )
        db.session.commit()

    login = client.post(
        "/auth/login",
        data={"username": "vedomost-admin", "password": "secret123"},
        follow_redirects=False,
    )
    assert login.status_code in (302, 303)

    page = client.get("/admin/grades/class/5Д?period_number=2")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Сводная ведомость" in html
    assert "summary_vedomost" in html

    downloaded = client.get("/admin/grades/class/5Д/download-vedomost")
    assert downloaded.status_code == 200
    assert downloaded.data[:2] == b"PK"
    ws = load_workbook(BytesIO(downloaded.data)).active
    assert ws["B11"].value == "Ли Арина"
    assert ws["L12"].value == 4
    assert ws["V24"].value == "Классный Руководитель"
