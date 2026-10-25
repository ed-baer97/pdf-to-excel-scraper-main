"""Сборка Excel-отчёта из JSON по шаблону: формативное оценивание, СОР, СОЧ, лист «Оценки»."""
import argparse
import json
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange


def _sanitize_filename(s: str) -> str:
    """Очищает строку для имени файла: убирает недопустимые символы Windows."""
    s = " ".join((s or "").split()).strip()
    s = re.sub(r"[<>:\"/\\\\|?*]+", "_", s)
    s = s.strip(" .")
    return s or "report"


def _load_json(path: Path):
    """Читает JSON-файл в UTF-8 и возвращает объект Python."""
    return json.loads(path.read_text(encoding="utf-8"))


def _points_by_section(points: dict, quarter_num: int) -> dict[int, str]:
    """Из словаря баллов по ключам `chetvert_{q}_razdel_{k}_...` берёт значения по номеру раздела k."""
    out: dict[int, str] = {}
    prefix = f"chetvert_{quarter_num}_razdel_"
    for pid, val in (points or {}).items():
        if not isinstance(pid, str) or not pid.startswith(prefix):
            continue
        parts = pid.split("_")
        # chetvert_{q}_razdel_{k}_{row}
        if len(parts) >= 5 and parts[2] == "razdel":
            try:
                sec = int(parts[3])
            except Exception:
                continue
            out[sec] = val
    return out


# Литера класса — целое слово из одной буквы, а не первая буква длинного названия.
_CLASS_LITER = r"[A-Za-zА-ЯЁӘҒҚҢӨҰҮҺа-яёәғқңөұүһ]"


def _parse_class_liter(class_text: str) -> str:
    """Преобразует текст класса в вид для листа.

    Одна литера целиком: «5 «В»» -> «5В».
    Словесное название не режется по первой букве:
    «6 «Дистанционное обучение»» остаётся «6 Дистанционное обучение».
    """
    s = (class_text or "").replace("«", " ").replace("»", " ")
    s = " ".join(s.split()).strip()
    if not s:
        return ""
    m = re.fullmatch(rf"(\d+)\s*({_CLASS_LITER})?", s)
    if not m:
        return s
    num = m.group(1)
    lit = (m.group(2) or "").upper()
    return f"{num}{lit}"


def _to_number(val):
    """Приводит значение к float (проценты, запятые, пустые строки → 0)."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace("%", "").replace(",", ".")
    try:
        return float(s) if s else 0.0
    except Exception:
        return 0.0


def _extract_max_points_from_criteria_html(criteria_html: Path, quarter_num: int) -> dict[int, int]:
    """Запасной вариант: парсит макс. баллы по разделам из сохранённого criteria.html (атрибуты max)."""
    if not criteria_html.exists():
        return {}
    html = criteria_html.read_text(encoding="utf-8", errors="ignore")
    out: dict[int, int] = {}
    pattern = re.compile(rf'id="chetvert_{quarter_num}_razdel_(\d+)_max"[^>]*\svalue="([^"]+)"', re.IGNORECASE)
    for sec_s, val_s in pattern.findall(html):
        try:
            sec = int(sec_s)
            val = int(float(val_s.strip()))
            out[sec] = val
        except Exception:
            continue
    return out


def _clear_rows(ws, start_row: int, end_row: int, cols: list[str]):
    """Обнуляет ячейки в диапазоне строк по указанным колонкам."""
    for r in range(start_row, end_row + 1):
        for col in cols:
            ws[f"{col}{r}"].value = None


# Строка заголовка таблицы — 7, первая строка ученика — 8.
# Ниже списка идёт строка-подвал (ссылки на итоги) и блок «Всего "5"…».
_STUDENT_START = 8
_STUDENT_COLS = tuple("ABCDEFGHIJKLMN")
_CELL_REF = re.compile(r"(\$?[A-Z]{1,3})(\$?)(\d+)")


def _summary_kind(value) -> str | None:
    """Метка итогового блока: '5'/'4'/'3'/'2' или 'total' для «Всего учеников»."""
    if not isinstance(value, str) or "всего" not in value.lower():
        return None
    text = value.replace("«", '"').replace("»", '"').replace("“", '"').replace("”", '"')
    text = " ".join(text.split()).lower()
    if "ученик" in text:
        return "total"
    for digit in ("5", "4", "3", "2"):
        if digit in text:
            return digit
    return None


def summary_rows(ws) -> dict[str, int]:
    """Строки блока итогов по подписям в колонке D. Ключи: 5, 4, 3, 2, total."""
    found: dict[str, int] = {}
    for row in ws.iter_rows(min_row=1, max_col=9, max_row=ws.max_row or 1):
        for cell in row:
            kind = _summary_kind(cell.value)
            if kind and kind not in found:
                found[kind] = cell.row
    return found


def student_last_row(ws) -> int:
    """Последняя строка ученика: на две выше «Всего "5"» (между ними строка-подвал)."""
    anchor = summary_rows(ws).get("5")
    if anchor and anchor - 2 >= _STUDENT_START:
        return anchor - 2
    return 39


def _retarget_formula(formula: str, src_row: int, dst_row: int) -> str:
    """В формуле заменяет ссылки на src_row ссылками на dst_row."""
    def repl(match: re.Match) -> str:
        col, absolute, row_s = match.group(1), match.group(2), match.group(3)
        if int(row_s) == src_row:
            return f"{col}{absolute}{dst_row}"
        return match.group(0)

    return _CELL_REF.sub(repl, formula)


def _copy_student_row(ws, src: int, dst: int) -> None:
    """Копирует оформление и формулы строки ученика, без ФИО и баллов."""
    for col in _STUDENT_COLS:
        src_cell = ws[f"{col}{src}"]
        dst_cell = ws[f"{col}{dst}"]
        if getattr(src_cell, "has_style", False):
            dst_cell.font = copy(src_cell.font)
            dst_cell.border = copy(src_cell.border)
            dst_cell.fill = copy(src_cell.fill)
            dst_cell.number_format = src_cell.number_format
            dst_cell.alignment = copy(src_cell.alignment)
        value = src_cell.value
        if isinstance(value, str) and value.startswith("="):
            dst_cell.value = _retarget_formula(value, src, dst)
        else:
            dst_cell.value = None
    height = ws.row_dimensions[src].height
    if height is not None:
        ws.row_dimensions[dst].height = height


def _unmerge_overlapping(ws, min_row: int, max_row: int) -> None:
    """Снимает объединения, которые пересекают удаляемые строки."""
    doomed = [
        str(merged)
        for merged in ws.merged_cells.ranges
        if merged.max_row >= min_row and merged.min_row <= max_row
    ]
    for ref in doomed:
        ws.unmerge_cells(ref)


def _shift_merges(ws, pivot_row: int, offset: int) -> None:
    """Сдвигает объединения ниже вставленных или удалённых строк.

    openpyxl переносит ячейки, но диапазоны merge оставляет на старых номерах.
    Снимать их через unmerge_cells нельзя: ячеек по старым координатам уже нет.
    """
    if offset == 0:
        return
    original = list(ws.merged_cells.ranges)
    for merged in original:
        ws.merged_cells.remove(merged)
    for merged in original:
        min_row, max_row = merged.min_row, merged.max_row
        if max_row < pivot_row:
            pass
        elif min_row >= pivot_row:
            min_row += offset
            max_row += offset
        else:
            max_row += offset
        if min_row < 1 or min_row > max_row:
            continue
        ws.merged_cells.add(CellRange(
            min_col=merged.min_col,
            min_row=min_row,
            max_col=merged.max_col,
            max_row=max_row,
        ))


def _refresh_summary_formulas(ws, last_row: int) -> dict[str, int]:
    """Перенацеливает формулы итогов на фактический диапазон учеников."""
    rows = summary_rows(ws)
    r5, r4, r3, r2 = rows.get("5"), rows.get("4"), rows.get("3"), rows.get("2")
    rtotal = rows.get("total")
    if not all((r5, r4, r3, r2, rtotal)):
        return rows
    start = _STUDENT_START
    ws[f"F{r5}"] = f"=SUM(F{start}:F{last_row})"
    ws[f"G{r4}"] = f"=SUM(G{start}:G{last_row})"
    ws[f"H{r3}"] = f"=SUM(H{start}:H{last_row})"
    ws[f"I{r2}"] = f"=SUM(I{start}:I{last_row})"
    ws[f"I{rtotal}"] = f"=F{r5}+G{r4}+H{r3}+I{r2}"
    ws[f"C{r5}"] = f"=(F{r5}+G{r4})/I{rtotal}"
    ws[f"C{r4}"] = f"=(F{r5}+G{r4}+H{r3})/I{rtotal}"
    link = r5 - 1
    ws[f"L{link}"] = f"=F{r5}"
    ws[f"M{link}"] = f"=G{r4}+H{r3}"
    ws[f"N{link}"] = f"=I{r2}"
    ws["J8"] = f"=C{r5}*100"
    ws["K8"] = f"=C{r4}*100"
    return rows


def resize_student_block(ws, count: int) -> int:
    """Подгоняет число строк учеников под `count`, не трогая шапку и блок итогов.

    Лишние строки удаляются, недостающие вставляются копией последней строки
    ученика (границы и формулы). Блок «Всего» остаётся сразу под списком.
    Возвращает номер последней строки ученика.
    """
    count = max(int(count or 0), 0)
    rows = summary_rows(ws)
    anchor = rows.get("5")
    if not anchor or anchor - 2 < _STUDENT_START:
        last = _STUDENT_START + max(count, 1) - 1
        return last

    link_row = anchor - 1
    last = link_row - 1
    capacity = last - _STUDENT_START + 1
    if count > capacity:
        extra = count - capacity
        ws.insert_rows(link_row, extra)
        _shift_merges(ws, link_row, extra)
        for offset in range(extra):
            _copy_student_row(ws, last, last + 1 + offset)
        last += extra
    elif count < capacity:
        keep = max(count, 1)
        first_delete = _STUDENT_START + keep
        if first_delete <= last:
            _unmerge_overlapping(ws, first_delete, last)
            deleted = last - first_delete + 1
            ws.delete_rows(first_delete, deleted)
            _shift_merges(ws, first_delete, -deleted)
            last = _STUDENT_START + keep - 1

    _refresh_summary_formulas(ws, last)
    for index, row in enumerate(range(_STUDENT_START, last + 1), start=1):
        ws[f"A{row}"].value = index
    return last


def _fill_template_page(
    ws,
    *,
    organization_name: str | None,
    class_liter: str | None,
    teacher_fio: str | None,
    page_title: str,
    fio_list: list[str],
    result_list: list,
    max_value,
    mode: str,
):
    """Заполняет один лист-копию шаблона: заголовок, ФИО, результат/проценты, флаги по уровням.

    mode: «percent» — значения как проценты 0–100; «points» — доля от max_value; «grade» — оценки 2–5.
    """
    # Header cells (as in example)
    if organization_name:
        ws["B1"].value = organization_name
    if class_liter:
        ws["C3"].value = class_liter
    ws["C4"].value = len(fio_list)
    if teacher_fio:
        ws["C5"].value = teacher_fio
    ws["C6"].value = page_title

    start_row = _STUDENT_START
    last_row = resize_student_block(ws, len(fio_list))

    # Ensure the template column header for max points is present (your note: D7 is the column name).
    if mode == "points":
        ws["D7"].value = "Макс"

    # Fill rows
    count_5 = count_4 = count_3 = count_2 = 0
    # If you want max shown "under D7" regardless of row, prefill the student block with the same max.
    if mode == "points" and max_value not in (None, "") and fio_list:
        mv = _to_number(max_value)
        mv = int(mv) if mv.is_integer() else mv
        for r in range(start_row, last_row + 1):
            ws[f"D{r}"].value = mv

    for i, fio in enumerate(fio_list):
        r = start_row + i
        if r > last_row:
            break
        ws[f"B{r}"].value = fio

        res = result_list[i] if i < len(result_list) else ""
        ws[f"C{r}"].value = res

        # D - max points already prefilled for points mode; clear for non-points.
        if mode != "points":
            ws[f"D{r}"].value = None

        # Compute percent and grade flags similar to example template logic
        percent = 0.0
        if mode == "percent":
            percent = _to_number(res)
        elif mode == "points":
            rv = _to_number(res)
            mv = _to_number(max_value)
            percent = (rv / mv * 100.0) if mv > 0 else 0.0
        elif mode == "grade":
            g = int(_to_number(res))
            # For grade mode, place flags directly (percent is not meaningful)
            ws[f"E{r}"].value = None
            ws[f"F{r}"].value = 1 if g == 5 else 0
            ws[f"G{r}"].value = 1 if g == 4 else 0
            ws[f"H{r}"].value = 1 if g == 3 else 0
            ws[f"I{r}"].value = 1 if g == 2 else 0
            if g == 5:
                count_5 += 1
                ws[f"L{r}"].value = fio
                ws[f"M{r}"].value = ""
                ws[f"N{r}"].value = ""
            elif g in (3, 4):
                if g == 4:
                    count_4 += 1
                else:
                    count_3 += 1
                ws[f"L{r}"].value = ""
                ws[f"M{r}"].value = fio
                ws[f"N{r}"].value = ""
            elif g == 2:
                count_2 += 1
                ws[f"L{r}"].value = ""
                ws[f"M{r}"].value = ""
                ws[f"N{r}"].value = fio
            else:
                ws[f"L{r}"].value = ""
                ws[f"M{r}"].value = ""
                ws[f"N{r}"].value = ""
            continue

        ws[f"E{r}"].value = round(percent, 2)
        is5 = percent >= 85
        is4 = 65 <= percent < 85
        is3 = 40 <= percent < 65
        is2 = percent < 40

        ws[f"F{r}"].value = 1 if is5 else 0
        ws[f"G{r}"].value = 1 if is4 else 0
        ws[f"H{r}"].value = 1 if is3 else 0
        ws[f"I{r}"].value = 1 if is2 else 0

        if is5:
            count_5 += 1
            ws[f"L{r}"].value = fio
            ws[f"M{r}"].value = ""
            ws[f"N{r}"].value = ""
        elif is4 or is3:
            if is4:
                count_4 += 1
            else:
                count_3 += 1
            ws[f"L{r}"].value = ""
            ws[f"M{r}"].value = fio
            ws[f"N{r}"].value = ""
        else:
            count_2 += 1
            ws[f"L{r}"].value = ""
            ws[f"M{r}"].value = ""
            ws[f"N{r}"].value = fio

    # Пустой класс: единственная оставленная строка не должна показывать формулы деления на ноль.
    if not fio_list:
        _clear_rows(ws, start_row, last_row, ["B", "C", "D", "E", "F", "G", "H", "I", "L", "M", "N"])

    # Итоги стоят сразу под списком, а не на фиксированных строках 41–44.
    totals = summary_rows(ws)
    ws[f"F{totals.get('5', 41)}"].value = count_5
    ws[f"G{totals.get('4', 42)}"].value = count_4
    ws[f"H{totals.get('3', 43)}"].value = count_3
    ws[f"I{totals.get('2', 44)}"].value = count_2

    total_students = len(fio_list)
    if total_students > 0:
        quality = ((count_4 + count_5) / total_students) * 100.0
        success_rate = ((count_3 + count_4 + count_5) / total_students) * 100.0
    else:
        quality = 0.0
        success_rate = 0.0
    # In user's example they wrote these to J8 / K8; keep same.
    ws["J8"].value = round(quality, 2)
    ws["K8"].value = round(success_rate, 2)


def _fill_grades_page(
    ws,
    *,
    organization_name: str | None,
    class_liter: str | None,
    teacher_fio: str | None,
    fio_list: list[str],
    nums: list[int],
    grades: list,
    percents: list,
):
    """Заполняет упрощённый лист «Оценки»: №, ФИО, оценка, %, итоговые качество/успеваемость в J8/K8."""
    if organization_name:
        ws["B1"].value = organization_name
    if class_liter:
        ws["C3"].value = class_liter
    ws["C4"].value = len(fio_list)
    if teacher_fio:
        ws["C5"].value = teacher_fio
    ws["C6"].value = "Оценки"

    header_row = 7
    start_row = _STUDENT_START
    last_row = resize_student_block(ws, len(fio_list))

    # Header
    ws[f"A{header_row}"].value = "№"
    ws[f"B{header_row}"].value = "ФИО"
    ws[f"C{header_row}"].value = "Оценка"
    ws[f"D{header_row}"].value = "%"
    # Per user: E and F must be empty. Put labels back to J7/K7.
    ws[f"E{header_row}"].value = None
    ws[f"F{header_row}"].value = None
    ws[f"J{header_row}"].value = "кач-ва"
    ws[f"K{header_row}"].value = "успев"

    # Remove/clear everything from L onwards (and keep G-I empty too)
    for col in ["G", "H", "I", "L", "M", "N"]:
        ws[f"{col}{header_row}"].value = None

    count_5 = count_4 = count_3 = count_2 = 0
    qual_flags = []
    succ_flags = []

    for i, fio in enumerate(fio_list):
        r = start_row + i
        if r > last_row:
            break
        num = nums[i] if i < len(nums) else i + 1
        g = grades[i] if i < len(grades) else ""
        p = percents[i] if i < len(percents) else ""

        g_num = int(_to_number(g)) if str(g).strip() != "" else 0
        if g_num == 5:
            count_5 += 1
        elif g_num == 4:
            count_4 += 1
        elif g_num == 3:
            count_3 += 1
        elif g_num == 2:
            count_2 += 1

        ws[f"A{r}"].value = num
        ws[f"B{r}"].value = fio
        ws[f"C{r}"].value = g_num if g_num else g
        ws[f"D{r}"].value = _to_number(p) if str(p).strip() != "" else p
        # E and F must be empty (no per-student flags)
        ws[f"E{r}"].value = None
        ws[f"F{r}"].value = None

        qual = 1 if g_num in (4, 5) else 0
        succ = 1 if g_num in (3, 4, 5) else 0
        qual_flags.append(qual)
        succ_flags.append(succ)

        # Clear everything from L onwards ("Высокий" etc), and keep G-I empty.
        for col in ["G", "H", "I", "L", "M", "N"]:
            ws[f"{col}{r}"].value = None

    if not fio_list:
        _clear_rows(ws, start_row, last_row, ["A", "B", "C", "D", "E", "F", "G", "H", "I", "L", "M", "N"])

    # Итоги стоят сразу под списком, а не на фиксированных строках 41–44.
    totals = summary_rows(ws)
    ws[f"F{totals.get('5', 41)}"].value = count_5
    ws[f"G{totals.get('4', 42)}"].value = count_4
    ws[f"H{totals.get('3', 43)}"].value = count_3
    ws[f"I{totals.get('2', 44)}"].value = count_2

    total = len(fio_list)
    quality = (sum(qual_flags) / total * 100.0) if total else 0.0
    success_rate = (sum(succ_flags) / total * 100.0) if total else 0.0
    ws["J8"].value = round(quality, 2)
    ws["K8"].value = round(success_rate, 2)

    # Колонки «высокий/средний/низкий» на листе оценок не заполняются.
    # Подвал под списком (ссылки на итоги) не трогаем.
    for r in range(header_row, last_row + 1):
        for col in ["L", "M", "N"]:
            ws[f"{col}{r}"].value = None


def build_report(
    template_path: Path,
    students_path: Path,
    context_path: Path,
    out_dir: Path,
    max_points_path: Path | None = None,
    criteria_html_path: Path | None = None,
    org_name_path: Path | None = None,
) -> Path:
    """Собирает xlsx из шаблона и JSON учеников/контекста; добавляет листы по данным и сохраняет в out_dir."""
    students = _load_json(students_path)
    ctx = _load_json(context_path)
    max_points = {}
    if max_points_path and max_points_path.exists():
        try:
            mp = _load_json(max_points_path)
            if isinstance(mp, dict):
                # keys might come as strings in JSON
                max_points = {int(k): int(v) for k, v in mp.items()}
        except Exception:
            max_points = {}

    class_name = str(ctx.get("class", "") or "").strip()
    subject = str(ctx.get("subject", "") or "").strip()
    quarter_num = int(students[0]["quarter_num"]) if students else 0
    org_name = ctx.get("org_name")
    if (not org_name) and org_name_path and org_name_path.exists():
        org_name = org_name_path.read_text(encoding="utf-8", errors="ignore").strip() or None
    teacher_fio = ctx.get("profile_name")
    class_liter = _parse_class_liter(class_name)

    # Fallback: if max_points file is missing/empty, parse from criteria.html
    if (not max_points) and quarter_num and criteria_html_path:
        max_points = _extract_max_points_from_criteria_html(criteria_html_path, quarter_num)

    out_dir.mkdir(parents=True, exist_ok=True)
    out_name = _sanitize_filename(f"{class_name} {subject}".strip())
    out_path = out_dir / f"{out_name}.xlsx"

    wb = load_workbook(template_path)
    template_ws = None
    for name in wb.sheetnames:
        if "шаблон" in name.lower() or "template" in name.lower():
            template_ws = wb[name]
            break
    if template_ws is None:
        template_ws = wb.worksheets[0]

    def mk_sheet(name: str):
        ws = wb.copy_worksheet(template_ws)
        ws.title = name[:31]
        return ws

    # Sort by student number to keep stable ordering
    students_sorted = sorted(students, key=lambda s: int(s.get("num") or 0))
    fio_list = [s.get("fio", "") for s in students_sorted]

    # Decide which pages exist (based on JSON fields / section points)
    # "Формативная оценка" in the table is the numeric value in the 3rd column (we store it as `average`).
    # Max is always 10 (per user).
    has_formative = any((s.get("average") not in (None, "", "0", 0)) for s in students_sorted)
    has_grades = any((s.get("grade") not in (None, "", "0", 0)) for s in students_sorted)

    sec_present: set[int] = set()
    for s in students_sorted:
        sec_present |= set(_points_by_section(s.get("points") or {}, quarter_num).keys())

    has_quarter_header = bool(ctx.get("has_quarter_grade_header"))

    # Pages in requested order
    if has_formative:
        ws = mk_sheet("Формативное оценивание")
        formative_vals = [s.get("average") or "" for s in students_sorted]
        _fill_template_page(
            ws,
            organization_name=org_name,
            class_liter=class_liter,
            teacher_fio=teacher_fio,
            page_title="Формативное оценивание",
            fio_list=fio_list,
            result_list=formative_vals,
            max_value=10,
            mode="points",
        )

    for sec in [1, 2, 3]:
        if sec not in sec_present:
            continue
        ws = mk_sheet(f"СОр {sec}")
        vals = []
        for s in students_sorted:
            pts = _points_by_section(s.get("points") or {}, quarter_num)
            vals.append(pts.get(sec, ""))
        _fill_template_page(
            ws,
            organization_name=org_name,
            class_liter=class_liter,
            teacher_fio=teacher_fio,
            page_title=f"СОр {sec}",
            fio_list=fio_list,
            result_list=vals,
            max_value=max_points.get(sec),
            mode="points",
        )

    if has_quarter_header and 0 in sec_present:
        ws = mk_sheet("СОч")
        vals = []
        for s in students_sorted:
            pts = _points_by_section(s.get("points") or {}, quarter_num)
            vals.append(pts.get(0, ""))
        _fill_template_page(
            ws,
            organization_name=org_name,
            class_liter=class_liter,
            teacher_fio=teacher_fio,
            page_title="СОч",
            fio_list=fio_list,
            result_list=vals,
            max_value=max_points.get(0),
            mode="points",
        )

    if has_grades and has_quarter_header:
        ws = mk_sheet("Оценки")
        nums = [int(s.get("num") or 0) for s in students_sorted]
        fio_list2 = [s.get("fio", "") for s in students_sorted]
        grade_vals = [s.get("grade") or "" for s in students_sorted]
        percent_vals = [s.get("total_pct") or "" for s in students_sorted]
        _fill_grades_page(
            ws,
            organization_name=org_name,
            class_liter=class_liter,
            teacher_fio=teacher_fio,
            fio_list=fio_list2,
            nums=nums,
            grades=grade_vals,
            percents=percent_vals,
        )

    # Keep original template sheet, rename explicitly
    # Per user: remove template sheet from final output.
    if template_ws.title in wb.sheetnames:
        wb.remove(template_ws)

    wb.save(out_path)
    return out_path


def main() -> int:
    """Точка входа CLI: пути к шаблону, JSON и выходной каталог."""
    p = argparse.ArgumentParser()
    p.add_argument("--template", default="Шаблон.xlsx")
    p.add_argument("--students", default="out/mektep/criteria_students.json")
    p.add_argument("--context", default="out/mektep/criteria_context.json")
    p.add_argument("--maxpoints", default="out/mektep/criteria_max_points.json")
    p.add_argument("--criteriahtml", default="out/mektep/criteria.html")
    p.add_argument("--orgfile", default="out/mektep/org_name.txt")
    p.add_argument("--outdir", default="out/mektep/reports")
    args = p.parse_args()

    out_path = build_report(
        template_path=Path(args.template),
        students_path=Path(args.students),
        context_path=Path(args.context),
        out_dir=Path(args.outdir),
        max_points_path=Path(args.maxpoints) if args.maxpoints else None,
        criteria_html_path=Path(args.criteriahtml) if args.criteriahtml else None,
        org_name_path=Path(args.orgfile) if args.orgfile else None,
    )
    print(f"Saved report: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

