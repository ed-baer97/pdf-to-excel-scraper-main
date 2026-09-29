"""Сводная ведомость класса по бланку «жинақ тізімдемесі»."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from ....constants import kazakh_sort_key, normalize_subject_name
from ...year_grades import compute_year_grade_from_periods
from ..payload import report_grades_payload

TEMPLATE_PATH = Path(__file__).with_name("templates") / "svodnaya_vedomost.xlsx"

# Столбцы D–W бланка: предмет и в какую колонку отделения его класть.
# both — одна оценка пишется в обе одноимённые колонки (технология, физкультура).
# all — предмет на весь класс, одна колонка.
# kk / ru — казахское или русское отделение.
FIXED_SLOTS: tuple[tuple[str, str], ...] = (
    ("foreign", "kk"),
    ("kazakh", "ru"),
    ("foreign", "ru"),
    ("technology", "both"),
    ("modeling", "all"),
    ("informatics", "kk"),
    ("kazakh", "kk"),
    ("informatics", "ru"),
    ("math", "all"),
    ("class_hour", "all"),
    ("global", "all"),
    ("science", "all"),
    ("russian", "all"),
    ("russian_lit", "all"),
    ("history_kz", "all"),
    ("history_world", "all"),
    ("pe", "both"),
    ("pe", "both"),
    ("technology", "both"),
    ("music", "all"),
)

FIRST_STUDENT_ROW = 11
ROWS_PER_STUDENT = 7
ELECTIVE_COLUMNS = (24, 25, 26, 27, 28)  # X–AB
QUARTER_ROWS = (1, 2, 3, 4)
PERIOD_LABELS = (
    "1 тоқсан/четверть",
    "2 тоқсан/четверть \nІ жартыжылдық полугодие",
    "3 тоқсан/четверть",
    "4 тоқсан/четверть \nІІ жартыжылдық полугодие",
    "Жылдық бағасы \nГодовая оценка",
    "Емтихан бағасы Экзаменационная оценка ",
    "Қорытынды баға \nИтоговая оценка ",
)

_SUBGROUP_RE = re.compile(r"\((\d+)\)\s*$")
_CLASS_GRADE_RE = re.compile(r"^(\d+)")
_PASS_MARKS = {"зч", "зач", "зачет", "зачёт", "зачтено"}
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FONT = Font(name="Arial", size=10, color="000000")
_FONT_BOLD = Font(name="Arial", size=10, bold=True, color="000000")
_FONT_CAPTION = Font(name="Times New Roman", size=8, color="000000")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
_BOTTOM = Border(bottom=_THIN)


def instruction_track(subgroup: int | None) -> str:
    """Колонка языка в бланке 5д.

    Подгруппа (1) — русское отделение: шетел тілі в F, қазақ тілі в E,
    информатика в K. Подгруппа (2) и предмет без номера — казахское: D, J, I.
    """
    if subgroup == 1:
        return "ru"
    return "kk"


def classify_subject(canonical: str) -> str | None:
    """Слот бланка или None, если предмет уходит в «по выбору»."""
    name = canonical.casefold().replace("ё", "е")
    if "русск" in name and "литерат" in name:
        return "russian_lit"
    if "литературн" in name and "чтен" in name:
        return "russian_lit"
    if name == "русский язык" or name.endswith(" русский язык"):
        return "russian"
    if "казахск" in name or "қазақ" in name:
        return "kazakh"
    if any(token in name for token in ("иностран", "английск", "немецк", "французск", "шетел")):
        return "foreign"
    if "информат" in name:
        return "informatics"
    if "технолог" in name:
        return "technology"
    if "физкультур" in name or ("физическ" in name and "культур" in name) or "дене шыны" in name:
        return "pe"
    if name == "математика":
        return "math"
    if "естествознан" in name or "жаратылыстану" in name or "познание мира" in name:
        return "science"
    if "истори" in name and "казахстан" in name:
        return "history_kz"
    if "всемирн" in name or "дүниежүзі" in name:
        return "history_world"
    if "музык" in name:
        return "music"
    if ("классн" in name and "час" in name) or "сынып сағат" in name:
        return "class_hour"
    if "глобальн" in name or "жаһандық" in name:
        return "global"
    if "модельдеу" in name or "моделирован" in name:
        return "modeling"
    return None


def format_grade(raw: Any) -> int | str | None:
    """Число 2–5, «зч», «осв» или None, если ячейку оставляем пустой."""
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        if isinstance(raw, float) and raw != int(raw):
            return None
        return int(raw)
    text = str(raw).strip()
    if not text or text in {"—", "-", "–"}:
        return None
    low = text.casefold().replace("ё", "е")
    if low in _PASS_MARKS or low.startswith("зачет") or low.startswith("зачт"):
        return "зч"
    if low.startswith("осв") or low.startswith("освобожд"):
        return "осв"
    try:
        number = float(low.replace(",", "."))
    except ValueError:
        return text
    if number == int(number):
        return int(number)
    return None


def columns_for(slot: str, track: str) -> list[int]:
    """Номера столбцов Excel (1-based) для слота и отделения."""
    columns: list[int] = []
    for index, (name, mode) in enumerate(FIXED_SLOTS):
        if name != slot:
            continue
        if mode in ("both", "all") or mode == track:
            columns.append(4 + index)
    return columns


def _subgroup_number(raw_name: str) -> int | None:
    match = _SUBGROUP_RE.search((raw_name or "").strip())
    if not match:
        return None
    return int(match.group(1))


def _prefer(old: int | str | None, new: int | str | None) -> int | str | None:
    if new is None:
        return old
    if old is None:
        return new
    if isinstance(old, int) and isinstance(new, int):
        return max(old, new)
    return new


def _period_key(period_type: str, period_number: int) -> int | str | None:
    if period_type == "quarter" and period_number in QUARTER_ROWS:
        return period_number
    if period_type == "semester" and period_number == 1:
        return 2
    if period_type == "semester" and period_number == 2:
        return 4
    if period_type == "year":
        return "year"
    if period_type == "final":
        return "final"
    if period_type == "exam":
        return "exam"
    return None


def _year_from_quarters(values: dict[int | str, int | str]) -> int | str | None:
    present: list[int | str] = []
    ints: dict[int, int] = {}
    for period in QUARTER_ROWS:
        value = values.get(period)
        if value is None:
            continue
        present.append(value)
        if isinstance(value, int):
            ints[period] = value
    if not present:
        return None
    if len(ints) == len(present):
        return compute_year_grade_from_periods({period: ints.get(period) for period in QUARTER_ROWS})
    if not ints and all(value == present[0] for value in present):
        return present[0]
    if not ints:
        return present[-1]
    return compute_year_grade_from_periods(ints)


def collect_vedomost_grid(
    reports: list[Any],
    school_id: int | None = None,
) -> tuple[list[str], dict[str, dict[int, dict[int | str, int | str]]], list[str]]:
    """Ученики, оценки по столбцам Excel и названия факультативов (до 5)."""
    semester_keys: set[tuple[str, str]] = set()
    events: list[tuple[str, str, str, int | str, int | str, str]] = []

    for report in reports:
        period = _period_key(
            str(getattr(report, "period_type", "") or ""),
            int(getattr(report, "period_number", 0) or 0),
        )
        if period is None:
            continue
        raw_subject = str(getattr(report, "subject_name", "") or "")
        canonical = normalize_subject_name(raw_subject, school_id)
        slot = classify_subject(canonical)
        track = instruction_track(_subgroup_number(raw_subject))
        kind = slot or f"elective:{canonical}"
        source = str(getattr(report, "period_type", "") or "")
        if source == "semester":
            semester_keys.add((kind, track))
        payload = report_grades_payload(report) or {}
        for student in payload.get("students") or []:
            name = str((student or {}).get("name") or "").strip()
            grade = format_grade((student or {}).get("grade"))
            if not name or grade is None:
                continue
            events.append((name, kind, track, period, grade, source))

    cells: dict[str, dict[int, dict[int | str, int | str]]] = {}
    elective_names: list[str] = []
    seen_electives: set[str] = set()

    for name, kind, track, period, grade, source in events:
        if source == "quarter" and period in (1, 3) and (kind, track) in semester_keys:
            continue
        if kind.startswith("elective:"):
            title = kind.split(":", 1)[1]
            if title not in seen_electives:
                seen_electives.add(title)
                elective_names.append(title)
        columns = (
            columns_for(kind, track)
            if not kind.startswith("elective:")
            else []
        )
        bucket = cells.setdefault(name, {})
        if columns:
            for column in columns:
                period_map = bucket.setdefault(column, {})
                period_map[period] = _prefer(period_map.get(period), grade)

    elective_names = sorted(elective_names, key=kazakh_sort_key)[: len(ELECTIVE_COLUMNS)]
    elective_column = {
        title: ELECTIVE_COLUMNS[index] for index, title in enumerate(elective_names)
    }
    for name, kind, track, period, grade, source in events:
        if not kind.startswith("elective:"):
            continue
        if source == "quarter" and period in (1, 3) and (kind, track) in semester_keys:
            continue
        title = kind.split(":", 1)[1]
        column = elective_column.get(title)
        if column is None:
            continue
        period_map = cells.setdefault(name, {}).setdefault(column, {})
        period_map[period] = _prefer(period_map.get(period), grade)

    for period_map in (column_map for student in cells.values() for column_map in student.values()):
        if "year" not in period_map:
            year_grade = _year_from_quarters(period_map)
            if year_grade is not None:
                period_map["year"] = year_grade

    cells = {name: columns for name, columns in cells.items() if columns}
    students = sorted(cells, key=kazakh_sort_key)
    return students, cells, elective_names


def class_heading(class_name: str, academic_year: int) -> str:
    match = _CLASS_GRADE_RE.match((class_name or "").strip())
    label = match.group(1) if match else (class_name or "").strip()
    return f"{label} сынып / класс  {academic_year}  оқу жылы / учебный год"


def _apply_border(cell, *, alignment: Alignment, font: Font | None = None, value: Any = None) -> None:
    if value is not None:
        cell.value = value
    cell.font = font or _FONT
    cell.alignment = alignment
    cell.border = _BORDER


def _write_students(ws, students: list[str], cells: dict, elective_names: list[str]) -> int:
    for index, title in enumerate(elective_names):
        header = ws.cell(10, ELECTIVE_COLUMNS[index], value=title)
        header.font = _FONT
        header.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True, textRotation=90
        )
        header.border = _BORDER

    grade_columns = list(range(4, 24)) + list(ELECTIVE_COLUMNS)
    period_keys = (1, 2, 3, 4, "year", "exam", "final")
    for student_index, name in enumerate(students):
        start = FIRST_STUDENT_ROW + student_index * ROWS_PER_STUDENT
        end = start + ROWS_PER_STUDENT - 1
        _apply_border(ws.cell(start, 1), alignment=_CENTER, value=student_index + 1)
        _apply_border(ws.cell(start, 2), alignment=_CENTER, value=name)
        for row in range(start, end + 1):
            ws.row_dimensions[row].height = 23.85
            if row != start:
                _apply_border(ws.cell(row, 1), alignment=_CENTER)
                _apply_border(ws.cell(row, 2), alignment=_CENTER)
            label = PERIOD_LABELS[row - start]
            _apply_border(ws.cell(row, 3), alignment=_LEFT, value=label)
            period_key = period_keys[row - start]
            student_cells = cells.get(name, {})
            for column in grade_columns:
                grade = student_cells.get(column, {}).get(period_key)
                _apply_border(
                    ws.cell(row, column),
                    alignment=_CENTER,
                    value=" " if grade is None else grade,
                )
        ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
        ws.merge_cells(start_row=start, start_column=2, end_row=end, end_column=2)
    if not students:
        return 10
    return FIRST_STUDENT_ROW + len(students) * ROWS_PER_STUDENT - 1


def _write_signature(ws, start_row: int, director: str, class_teacher: str) -> int:
    labels = (
        (0, "Орта білім беру ұйымының директоры", True),
        (1, "Директор организации среднего образования", False),
        (2, "Мөрдің орны", False),
        (3, "Место печати", False),
        (4, "Сынып жетекшісі", True),
        (5, "Классный руководитель ", False),
    )
    for offset, text, bold in labels:
        row = start_row + offset
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=21)
        cell = ws.cell(row, 2, value=text)
        cell.font = _FONT_BOLD if bold else _FONT
        cell.alignment = _RIGHT
        ws.row_dimensions[row].height = 15
    ws.row_dimensions[start_row + 6].height = 15

    for name_row, person in ((start_row, director), (start_row + 4, class_teacher)):
        for row in (name_row, name_row + 1):
            for column in range(22, 27):
                cell = ws.cell(row, column)
                cell.border = _BOTTOM
                cell.font = _FONT
        name_cell = ws.cell(name_row, 22, value=person or None)
        name_cell.font = _FONT
        name_cell.alignment = _CENTER
        ws.merge_cells(start_row=name_row, start_column=22, end_row=name_row + 1, end_column=26)

    for slash_row in (start_row + 1, start_row + 5):
        slash = ws.cell(slash_row, 27, value="/")
        slash.font = _FONT
        slash.alignment = Alignment(horizontal="left", vertical="center")
        slash.border = _BOTTOM

    for caption_row in (start_row + 2, start_row + 6):
        ws.merge_cells(start_row=caption_row, start_column=27, end_row=caption_row, end_column=28)
        caption = ws.cell(caption_row, 27, value="(қолы/подпись)")
        caption.font = _FONT_CAPTION
        caption.alignment = _CENTER
    return start_row + 6


def build_summary_vedomost_workbook(
    *,
    class_name: str,
    academic_year: int,
    school_name: str,
    class_teacher: str,
    reports: list[Any],
    school_id: int | None = None,
    director: str = "",
) -> tuple[BytesIO, str]:
    """Собирает xlsx по шаблону бланка и данным отчётов класса."""
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= FIRST_STUDENT_ROW:
            ws.unmerge_cells(str(merged))
    if ws.max_row and ws.max_row >= FIRST_STUDENT_ROW:
        for row in ws.iter_rows(min_row=FIRST_STUDENT_ROW, max_row=ws.max_row, max_col=28):
            for cell in row:
                cell.value = None

    if school_name:
        ws.cell(5, 1).value = school_name
    ws.cell(8, 1).value = class_heading(class_name, academic_year)

    students, cells, electives = collect_vedomost_grid(reports, school_id)
    last_student_row = _write_students(ws, students, cells, electives)
    signature_start = last_student_row + 3
    last_row = _write_signature(ws, signature_start, director, class_teacher)
    ws.print_area = f"A1:{get_column_letter(28)}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.scale = 79
    ws.page_setup.paperSize = 9
    if ws.sheet_properties.pageSetUpPr is not None:
        ws.sheet_properties.pageSetUpPr.fitToPage = False

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    safe_class = (class_name or "class").replace(" ", "_")
    filename = f"Сводная_ведомость_{safe_class}_{academic_year}.xlsx"
    return output, filename


def export_class_summary_vedomost(
    school_id: int,
    class_name: str,
    academic_year: int | None = None,
) -> tuple[BytesIO, str]:
    """Ведомость класса из GradeReport за учебный год."""
    from ....extensions import db
    from ....models import Class, GradeReport, School
    from ...academic_year import resolve_academic_year

    year = resolve_academic_year(academic_year)
    school = db.session.get(School, school_id)
    school_name = school.name if school else ""
    school_class = Class.query.filter_by(school_id=school_id, name=class_name).first()
    teacher = ""
    if school_class is not None and school_class.class_teacher is not None:
        teacher = (
            school_class.class_teacher.full_name or school_class.class_teacher.username or ""
        ).strip()
    reports = GradeReport.query.filter_by(
        school_id=school_id,
        class_name=class_name,
        academic_year=year,
    ).all()
    return build_summary_vedomost_workbook(
        class_name=class_name,
        academic_year=year,
        school_name=school_name,
        class_teacher=teacher,
        reports=reports,
        school_id=school_id,
    )
